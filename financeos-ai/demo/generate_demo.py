"""Reproducible fictional Northstar Cloud P&L. No real company data."""
from pathlib import Path
import csv
import math
import random
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def generate():
    rng=random.Random(42)
    rows=[]
    # Each account is attributable to one department; categories remain consistent.
    accounts=[('Commercial','Subscription revenue','revenue',680000),('Commercial','Services revenue','revenue',125000),
              ('Delivery','Cloud infrastructure','cogs',195000),('Delivery','Service delivery','cogs',112000),
              ('Product','Product payroll','opex',153000),('Commercial','Sales & marketing','opex',108000),
              ('Corporate','Corporate overhead','opex',59000)]
    for m in range(1,10):
        for dept,account,category,base in accounts:
            budget=round(base*(1+.022*(m-1)),2)
            factor=1+rng.uniform(-.025,.025)
            if m>=7 and account=='Subscription revenue':factor-=.065
            if m>=7 and account=='Cloud infrastructure':factor+=.12
            if m==9 and account=='Sales & marketing':factor+=.21
            actual=round(budget*factor,2)
            for scenario,amount in [('actual',actual),('budget',budget)]:
                rows.append([f'2026-{m:02d}',dept,account,category,scenario,f'{amount:.2f}','USD'])
    root=Path(__file__).parent
    headers=['period','department','account','category','scenario','amount','currency']
    with (root/'northstar_synthetic.csv').open('w',newline='') as f:
        w=csv.writer(f);w.writerow(headers);w.writerows(rows)
    with (root/'upload_template.csv').open('w',newline='') as f:
        w=csv.writer(f);w.writerow(headers);w.writerows(rows[:14])
    dirty=[rows[0],rows[0],['2026-13','Commercial','Subscription revenue','revenue','actual','not a number','USD'],
           ['2026-01','Corporate','Office','mystery','budget','500.00','EUR']]
    with (root/'dirty_upload_demo.csv').open('w',newline='') as f:
        w=csv.writer(f);w.writerow(headers);w.writerows(dirty)
    wb=Workbook();ws=wb.active;ws.title='Synthetic P&L';ws.append(headers)
    for r in rows:ws.append(r[:-2]+[float(r[-2]),r[-1]])
    for c in ws[1]:c.font=Font(color='FFFFFF',bold=True);c.fill=PatternFill('solid',fgColor='185B46')
    for col in 'ABCDEFG':ws.column_dimensions[col].width=26
    ws.freeze_panes='A2';wb.save(root/'northstar_synthetic.xlsx')
    (root/'README.md').write_text('''# Synthetic demonstration data
Northstar Cloud is fictional. All figures are generated with a fixed random seed (42), in USD, January–September 2026. The demo contains 126 monthly account/scenario rows across four departments and seven accounts.

Late-period subscription revenue is below its synthetic budget, cloud infrastructure expense is above budget, and September marketing expense is above budget. These are deliberately injected numerical patterns, not evidence of real pricing, customer, vendor or campaign events. The app never promotes these patterns to verified business causes.

Run `python demo/generate_demo.py` from the project root to recreate CSV/XLSX. The dirty demo intentionally contains a duplicate and invalid rows to show the validation gate. The small CSV template contains one complete month. Forecasting needs at least three months.
''')

if __name__=='__main__':generate()
