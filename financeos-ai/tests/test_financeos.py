import io
import json
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

import pandas as pd
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook, Workbook
from financeos import data, analysis, agent
from financeos.main import app


@pytest.fixture
def client(tmp_path,monkeypatch):
    monkeypatch.setenv('FINANCEOS_DB',str(tmp_path/'test.db'))
    monkeypatch.delenv('OPENAI_API_KEY',raising=False)
    monkeypatch.delenv('OPENAI_MODEL',raising=False)
    with TestClient(app) as c:yield c


def dataset(client):return client.get('/api/datasets').json()[0]['id']


def frame():
    return pd.DataFrame([['2026-01','Sales','Sales','revenue','actual','110.00'],['2026-01','Sales','Sales','revenue','budget','100.00'],
        ['2026-01','Ops','Cloud','cogs','actual','40.00'],['2026-01','Ops','Cloud','cogs','budget','30.00'],
        ['2026-01','HQ','Payroll','opex','actual','20.00'],['2026-01','HQ','Payroll','opex','budget','25.00']],columns=data.REQUIRED)


def upload(client,df,preview=False):
    return client.post('/api/upload',params={'preview':str(preview).lower()},files={'file':('test.csv',df.to_csv(index=False).encode(),'text/csv')})


def test_seed_and_health(client):
    assert client.get('/').status_code==200
    assert client.get('/static/app.js').status_code==200
    assert client.get('/api/health').json()=={'status':'ok'}
    assert client.get('/api/config').json()['model_configured'] is False
    s=client.get(f'/api/datasets/{dataset(client)}/summary').json()
    assert s['dataset']['synthetic']==1 and s['row_count']==126
    assert s['complete']


def test_exact_financial_math_and_favorability(client):
    id=upload(client,frame()).json()['id']
    s=analysis.summary(id)
    assert s['totals']['actual']['ebitda']==50
    assert s['totals']['budget']['ebitda']==45
    assert sum(r['ebitda_impact'] for r in s['accounts'])==5
    payroll=next(r for r in s['accounts'] if r['label']=='Payroll')
    assert payroll['variance']==-5 and payroll['favorable'] and payroll['variance_pct']==-20
    assert next(k for k in s['kpis'] if k['name']=='ebitda')['variance']==5


def test_preview_does_not_write_and_errors_atomic(client):
    n=len(data.datasets())
    assert upload(client,frame(),True).json()['validation']['accepted_rows']==6
    assert len(data.datasets())==n
    df=frame();df.loc[0,'amount']='NaN'
    r=upload(client,df).json()
    assert r['id'] is None and r['validation']['errors'][0]['row']==2
    assert len(data.datasets())==n


@pytest.mark.parametrize('col,value',[('amount','inf'),('amount','1.001'),('amount','100000000001'),('amount','=SUM(2,3)'),('period','2026-13'),('scenario','forecast'),('category','cash'),('department','=CMD()'),('account','')])
def test_reject_bad_values(client,col,value):
    df=frame();df.loc[0,col]=value
    assert upload(client,df).json()['validation']['errors']


def test_cleaning_duplicates_and_conflicts(client):
    df=frame();df.loc[0,'amount']=' $1,100.00 ';df.loc[0,'scenario']=' ACTUAL '
    df=pd.concat([df,df.iloc[[0]]],ignore_index=True)
    r=upload(client,df).json()
    assert r['id'] and r['validation']['accepted_rows']==6
    assert any('duplicate' in w for w in r['validation']['warnings'])
    df.loc[len(df)-1,'amount']='99'
    assert upload(client,df).json()['id'] is None


def test_unpaired_no_fabricated_zero_variance(client):
    df=frame().iloc[1:]
    id=upload(client,df).json()['id'];s=analysis.summary(id)
    assert not s['complete'] and all(k['variance'] is None for k in s['kpis'])
    assert next(r for r in s['accounts'] if r['label']=='Sales')['variance'] is None
    r=client.post(f'/api/datasets/{id}/investigate',json={'question':'Why EBITDA gap?'}).json()
    assert 'Incomplete' in r['headline']


def test_zero_budget_signed_corrections(client):
    df=frame();df.loc[1,'amount']='0';df.loc[4,'amount']='(2.00)'
    id=upload(client,df).json()['id'];s=analysis.summary(id)
    sales=next(r for r in s['accounts'] if r['label']=='Sales')
    assert sales['variance_pct'] is None
    assert s['totals']['actual']['opex']==-2


def test_xlsx_roundtrip_and_formula_rejected(client):
    raw=(data.ROOT/'demo'/'northstar_synthetic.xlsx').read_bytes()
    r=client.post('/api/upload?preview=false',files={'file':('demo.xlsx',raw)}).json()
    assert r['id'] and r['validation']['accepted_rows']==126
    wb=Workbook();ws=wb.active;ws.append(data.REQUIRED);ws.append(['2026-01','HQ','Payroll','opex','actual','=1+2'])
    b=io.BytesIO();wb.save(b)
    r=client.post('/api/upload',files={'file':('formula.xlsx',b.getvalue())}).json()
    assert r['validation']['errors']


def test_category_mapping_and_currency(client):
    df=frame();df.loc[1,'category']='opex'
    assert upload(client,df).json()['validation']['errors']
    df=frame();df['currency']='EUR'
    assert len(upload(client,df).json()['validation']['errors'])==6


def test_filters_and_sql_injection(client):
    id=dataset(client)
    a=client.get(f'/api/datasets/{id}/summary',params={'start':'2026-09','end':'2026-09','department':'Commercial'}).json()
    assert a['row_count']==6 and len(a['months'])==1
    bad=client.get(f'/api/datasets/{id}/records',params={'department':"' OR 1=1 --"})
    assert bad.json()==[]
    assert client.get(f'/api/datasets/{id}/summary',params={'start':'2027-01'}).status_code==400
    assert client.get('/api/datasets/missing/summary').status_code==400


def test_agent_trace_links_and_cause_boundary(client):
    id=dataset(client)
    r=client.post(f'/api/datasets/{id}/investigate',json={'question':'Why is EBITDA below budget?','start':'2026-09','end':'2026-09'}).json()
    assert r['mode']=='offline automation' and len(r['trace'])==4
    ids={r['id'] for r in data.records(id,'2026-09','2026-09')}
    assert all(set(e['record_ids'])<=ids for e in r['evidence'])
    assert 'not verified business causes' in r['unknowns'][0]
    assert all('customer churn' not in c.lower() for c in r['claims'])
    restored=client.get(f"/api/results/investigations/{r['id']}").json()
    assert restored==r
    assert client.get(f'/api/datasets/{id}/history').json()['investigations']


def test_unsupported_and_model_missing(client):
    id=dataset(client)
    r=client.post(f'/api/datasets/{id}/investigate',json={'question':'What is the stock price?'}).json()
    assert r['unsupported'] and r['evidence']==[]
    assert client.post(f'/api/datasets/{id}/investigate',json={'question':'Why EBITDA?','use_model':True}).status_code==400


def test_model_plan_validated_no_arbitrary_text(monkeypatch):
    monkeypatch.setenv('OPENAI_API_KEY','unit-test-fake')
    monkeypatch.setenv('OPENAI_MODEL','configured-model')
    class FakeResponse(io.BytesIO):
        def __enter__(self):return self
        def __exit__(self,*args):self.close()
    response={'output':[{'content':[{'type':'output_text','text':'{"intent":"revenue"}'}]}]}
    with patch('urllib.request.urlopen',return_value=FakeResponse(json.dumps(response).encode())) as call:
        assert agent.plan('Revenue gap',True)[0]=='revenue'
        body=json.loads(call.call_args[0][0].data)
        assert body['store'] is False and 'text' in body and body['input']=='Revenue gap'
    with patch('urllib.request.urlopen',side_effect=RuntimeError('sensitive provider failure')):
        intent,mode,notice=agent.plan('Revenue gap',True)
        assert intent=='revenue' and 'sensitive' not in notice
    response['output'][0]['content'][0]['text']='{"intent":"DROP TABLE records"}'
    with patch('urllib.request.urlopen',return_value=FakeResponse(json.dumps(response).encode())):
        assert agent.plan('Revenue gap',True)[0]=='revenue'


def test_forecast_base_and_downside(client):
    id=dataset(client)
    base=client.post(f'/api/datasets/{id}/forecast',json={}).json()
    assert base['points'][0]['period']=='2026-10' and len(base['points'])==6
    assert all(p['base_ebitda']==p['ebitda'] for p in base['points'])
    assert base['backtest']['observations']==6
    assert base['baseline']['periods']==['2026-07','2026-08','2026-09']
    downside=client.post(f'/api/datasets/{id}/forecast',json={'growth':-1,'revenue_shock':-5,'cogs_shift':2,'opex_shift':5}).json()
    assert downside['points'][0]['ebitda']<base['points'][0]['ebitda']
    for p in downside['points']:
        assert abs(p['revenue']-p['cogs']-p['opex']-p['ebitda'])<.02
    assert client.get(f"/api/results/scenarios/{downside['id']}").json()==downside


def test_forecast_insufficient_missing_and_invalid(client):
    id=dataset(client)
    assert client.post(f'/api/datasets/{id}/forecast',json={'start':'2026-09'}).status_code==400
    assert client.post(f'/api/datasets/{id}/forecast',json={'horizon':25}).status_code==422
    rows=data.records(id)
    df=pd.DataFrame([{k:(r['cents']/100 if k=='amount' else r[k]) for k in data.REQUIRED} for r in rows if r['period']!='2026-04'])
    gapid=upload(client,df).json()['id']
    assert client.post(f'/api/datasets/{gapid}/forecast',json={}).status_code==400
    df=df[~((df.period=='2026-03')&(df.account=='Product payroll'))]
    unbalanced=upload(client,df).json()['id']
    assert client.post(f'/api/datasets/{unbalanced}/forecast',json={}).status_code==400


def test_reports_scope_xlsx_and_html_safety(client):
    id=dataset(client)
    inv=client.post(f'/api/datasets/{id}/investigate',json={'question':'Why EBITDA gap?'}).json()
    scen=client.post(f'/api/datasets/{id}/forecast',json={}).json()
    params={'investigation_id':inv['id'],'scenario_id':scen['id']}
    html=client.get(f'/api/datasets/{id}/report',params=params)
    assert html.status_code==200 and 'SYNTHETIC DEMO DATA' in html.text and 'Source evidence appendix' in html.text
    assert inv['id'] in html.text
    xlsx=client.get(f'/api/datasets/{id}/report',params={**params,'format':'xlsx'})
    wb=load_workbook(io.BytesIO(xlsx.content));assert 'Scenario' in wb.sheetnames and wb['Source rows'].max_row==127
    assert client.get(f'/api/datasets/{id}/report',params={**params,'start':'2026-09'}).status_code==400
    other=client.post('/api/demo').json()['id']
    assert client.get(f'/api/datasets/{other}/report',params=params).status_code==400
    assert 'variance' in client.get(f'/api/datasets/{id}/report?format=csv').text
    df=frame();df['department']='<script>alert(1)</script>'
    safeid=upload(client,df).json()['id']
    report=client.get(f'/api/datasets/{safeid}/report').text
    assert '<script>alert(1)</script>' not in report and '&lt;script&gt;' in report


def test_origin_host_security(client):
    assert client.post('/api/demo',headers={'Origin':'https://evil.example'}).status_code==403
    assert client.get('/api/health',headers={'Host':'evil.example'}).status_code==403


def test_durable_storage(client):
    id=dataset(client)
    old=data.metadata(id)
    with data.connection() as db:
        assert db.execute('SELECT COUNT(*) FROM records WHERE dataset_id=?',(id,)).fetchone()[0]==126
    data.init_db()
    assert data.metadata(id)['sha256']==old['sha256']


def test_duplicate_csv_headers_rejected(client):
    raw=b'period,department,account,category,scenario,amount,amount\n2026-01,HQ,Payroll,opex,actual,10,999\n'
    r=client.post('/api/upload',files={'file':('duplicate.csv',raw)})
    assert r.status_code==400 and 'Duplicate column' in r.json()['detail']


def test_cogs_intent_headline(client):
    id=dataset(client)
    r=client.post(f'/api/datasets/{id}/investigate',json={'question':'Investigate COGS against budget'}).json()
    assert r['intent']=='cogs' and r['headline'].startswith('COGS')
    assert all(e['category']=='cogs' for e in r['evidence'])


def test_month_series_distinguishes_missing_scenario(client):
    df=frame();df=df[df.scenario=='actual']
    id=upload(client,df).json()['id']
    s=analysis.summary(id)
    assert s['months'][0]['present']=={'actual':True,'budget':False}
    assert s['complete'] is False


@pytest.fixture
def hosted_client(tmp_path,monkeypatch):
    monkeypatch.setenv('FINANCEOS_DB',str(tmp_path/'hosted.db'))
    monkeypatch.setenv('FINANCEOS_MODE','hosted')
    monkeypatch.setenv('RENDER_EXTERNAL_HOSTNAME','financeos-test.onrender.com')
    monkeypatch.setenv('FINANCEOS_AUTH_USERNAME','demo')
    monkeypatch.setenv('FINANCEOS_AUTH_PASSWORD','test-only-password-1234567890')
    monkeypatch.setenv('FINANCEOS_EPHEMERAL','true')
    with TestClient(app,base_url='https://financeos-test.onrender.com') as c:
        yield c


def test_hosted_auth_protects_every_data_route(hosted_client):
    c=hosted_client
    assert c.get('/api/health').json()=={'status':'ok'}
    for route in ('/','/api/config','/api/datasets','/static/app.js','/docs','/api/template'):
        denied=c.get(route)
        assert denied.status_code==401
        assert 'Basic realm=' in denied.headers['www-authenticate']
    assert c.get('/api/datasets',auth=('demo','wrong')).status_code==401
    assert c.get('/api/datasets',headers={'Authorization':'Basic bad-data'}).status_code==401
    assert c.get('/api/datasets',headers={'Authorization':'Bearer arbitrary'}).status_code==401
    response=c.get('/api/datasets',auth=('demo','test-only-password-1234567890'))
    assert response.status_code==200
    assert response.headers['Cache-Control']=='no-store'
    assert response.headers['Strict-Transport-Security']=='max-age=31536000'


def test_hosted_workflow_and_origin(hosted_client):
    c=hosted_client
    c.auth=('demo','test-only-password-1234567890')
    config=c.get('/api/config').json()
    assert config['hosted'] and config['ephemeral']
    assert 'password' not in json.dumps(config).lower()
    id=c.get('/api/datasets').json()[0]['id']
    assert c.get(f'/api/datasets/{id}/summary').status_code==200
    assert c.post(f'/api/datasets/{id}/forecast',json={},headers={'Origin':'https://financeos-test.onrender.com'}).status_code==200
    assert c.post('/api/demo',headers={'Origin':'http://financeos-test.onrender.com'}).status_code==403
    assert c.post('/api/demo',headers={'Origin':'https://evil.example'}).status_code==403
    assert c.get('/api/health',headers={'Host':'evil.example'}).status_code==403
    assert c.post('/api/demo',headers={'Content-Length':str(12*1024*1024)}).status_code==413


def test_hosted_startup_fails_closed(monkeypatch):
    from financeos import hosting
    monkeypatch.setenv('FINANCEOS_MODE','hosted')
    monkeypatch.delenv('FINANCEOS_AUTH_PASSWORD',raising=False)
    with pytest.raises(RuntimeError,match='PASSWORD'):
        hosting.validate_startup()
    monkeypatch.setenv('FINANCEOS_AUTH_PASSWORD','test-only-password-1234567890')
    monkeypatch.delenv('RENDER_EXTERNAL_HOSTNAME',raising=False)
    monkeypatch.setenv('FINANCEOS_ALLOWED_HOSTS','*.example.com')
    with pytest.raises(RuntimeError,match='HOSTNAME'):
        hosting.validate_startup()
    monkeypatch.setenv('FINANCEOS_ALLOWED_HOSTS','financeos.example.com')
    hosting.validate_startup()
    assert '*.example.com' not in hosting.allowed_hosts()
