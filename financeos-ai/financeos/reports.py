import csv
import io
from html import escape
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from . import data


def safe_cell(value):
    if isinstance(value,str) and value.startswith(('=','+','-','@','\t','\r')):
        return "'"+value
    return value


def workbook(s, investigation=None, scenario=None):
    out=io.BytesIO()
    source=data.records(s['dataset']['id'],**s['scope'])
    frames={'Read me':[{'item':'Dataset','value':s['dataset']['id']},{'item':'Source','value':s['dataset']['name']},
        {'item':'Synthetic','value':bool(s['dataset']['synthetic'])},{'item':'Scope','value':str(s['scope'])},
        {'item':'Method','value':s['methodology']},{'item':'SHA256','value':s['dataset']['sha256']}],
        'KPI summary':s['kpis'],'Account variance':[{k:v for k,v in r.items() if k!='record_ids'} for r in s['accounts']],
        'Source rows':source,
        'Evidence map':[{'account':r['label'],'record_ids':','.join(map(str,r['record_ids']))} for r in s['accounts']],
        'Validation':[{'message':w} for w in s['dataset']['validation']['warnings']]}
    if investigation:
        frames['Investigation']=[{'statement':investigation['headline']}]+[{'statement':c} for c in investigation['claims']+investigation['unknowns']]
        frames['Investigation scope']=[{'key':k,'value':str(v)} for k,v in investigation['scope'].items()]
    if scenario:
        frames['Scenario']=scenario['points']
        frames['Assumptions']=[{'assumption':k,'value':v} for k,v in scenario['assumptions'].items()]
        frames['Forecast method']=[{'key':k,'value':str(v)} for k,v in scenario['baseline'].items()]+[{'key':'backtest','value':str(scenario['backtest'])}]+[{'key':'limitation','value':v} for v in scenario['limitations']]
    with pd.ExcelWriter(out,engine='openpyxl') as writer:
        for name,records in frames.items():
            df=pd.DataFrame(records)
            for col in df.columns:
                df[col]=df[col].map(safe_cell)
            df.to_excel(writer,sheet_name=name,index=False)
            ws=writer.sheets[name]
            ws.freeze_panes='A2'
            ws.auto_filter.ref=ws.dimensions
            for cell in ws[1]:
                cell.font=Font(color='FFFFFF',bold=True)
                cell.fill=PatternFill('solid',fgColor='153E35')
            for cells in ws.columns:
                width=min(65,max(15,max(len(str(c.value or '')) for c in cells)+2))
                ws.column_dimensions[cells[0].column_letter].width=width
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment=Alignment(vertical='top',wrap_text=True)
                    if isinstance(cell.value,float): cell.number_format='#,##0.00;[Red](#,##0.00)'
    return out.getvalue()


def csv_export(s):
    out=io.StringIO()
    fields=['label','category','actual','budget','variance','variance_pct','ebitda_impact','complete']
    w=csv.DictWriter(out,fieldnames=fields,extrasaction='ignore')
    w.writeheader()
    for r in s['accounts']:
        w.writerow({k:safe_cell(v) for k,v in r.items() if k in fields})
    return out.getvalue()


def html_report(s, investigation=None, scenario=None):
    def money(v): return 'Unavailable' if v is None else f'${v:,.2f}'
    def e(v): return escape(str(v))
    kpis=''.join(f"<div><small>{e(k['name'].replace('_',' ').upper())}</small><h2>{money(k['actual'])}</h2><p>Budget {money(k['budget'])} · Δ {money(k['variance'])}</p></div>" for k in s['kpis'])
    rows=''.join(f"<tr><td>{e(r['label'])}</td><td>{money(r['actual'])}</td><td>{money(r['budget'])}</td><td>{money(r['variance'])}</td><td>{money(r['ebitda_impact'])}</td><td>{e(', '.join(map(str,r['record_ids'])))}</td></tr>" for r in s['accounts'])
    detail=''
    if investigation:
        detail=f"<h2>Financial investigation</h2><p>{e(investigation['headline'])}</p><p>Run {e(investigation['id'])} · {e(investigation['mode'])}</p>"+''.join(f'<p>{e(c)}</p>' for c in investigation['claims']+investigation['unknowns'])
    forecast=''
    if scenario:
        forecast='<h2>Scenario outlook</h2><p>'+e(scenario['assumptions'])+'</p><p>Baseline: '+e(scenario['baseline'])+'</p><table><tr><th>Month</th><th>Revenue</th><th>Base EBITDA</th><th>Scenario EBITDA</th></tr>'
        forecast+=''.join(f"<tr><td>{e(p['period'])}</td><td>{money(p['revenue'])}</td><td>{money(p['base_ebitda'])}</td><td>{money(p['ebitda'])}</td></tr>" for p in scenario['points'])+'</table><p>'+e(scenario['backtest'])+'</p><p>'+e(' '.join(scenario['limitations']))+'</p>'
    warnings=''.join(f'<li>{e(w)}</li>' for w in s['dataset']['validation']['warnings'])
    source=data.records(s['dataset']['id'],**s['scope'])
    appendix=''.join(f"<tr><td>{r['id']}</td><td>{r['source_row']}</td><td>{e(r['period'])}</td><td>{e(r['department'])}</td><td>{e(r['account'])}</td><td>{e(r['scenario'])}</td><td>{money(r['cents']/100)}</td></tr>" for r in source)
    return f'''<!doctype html><html lang="en"><meta charset="utf-8"><title>FinanceOS • Management report</title>
    <style>body{{font:15px/1.5 system-ui;color:#17352d;max-width:1100px;margin:45px auto;padding:25px}}small{{letter-spacing:2px}}h1{{font-size:40px}}.kpis{{display:flex;gap:25px;flex-wrap:wrap}}table{{border-collapse:collapse;width:100%;font-size:12px}}td,th{{text-align:left;padding:10px;border-bottom:1px solid #ddd}}h2{{break-after:avoid}}tr{{break-inside:avoid}}.badge{{color:#795218;background:#fff2d8;padding:8px}}@media print{{button{{display:none}}body{{margin:0;padding:0}}}}</style>
    <button onclick="window.print()">Print / save as PDF</button><p>FINANCEOS AI / MANAGEMENT PACK</p><h1>Performance, with evidence.</h1>
    <p class="badge">{'SYNTHETIC DEMO DATA — fictional Northstar Cloud' if s['dataset']['synthetic'] else 'USER UPLOAD — verify source provenance'} · USD · {'Complete comparison coverage' if s['complete'] else 'INCOMPLETE COMPARISON COVERAGE'}</p>
    <p>Dataset {e(s['dataset']['id'])} · {e(s['dataset']['name'])} · Created {e(s['dataset']['created'])}<br>Scope: {e(s['scope'])}<br>Source SHA-256: {e(s['dataset']['sha256'])}</p><div class="kpis">{kpis}</div>
    {detail}<h2>Account variance & evidence index</h2><table><tr><th>Account</th><th>Actual</th><th>Budget</th><th>Variance</th><th>EBITDA impact</th><th>Record IDs</th></tr>{rows}</table>
    {forecast}<h2>Methodology & validation</h2><p>{e(s['methodology'])}</p><ul>{warnings}</ul><p>Arithmetic contributors do not establish business causes. Signed corrections retained. Scenario estimates are assumptions, not guarantees.</p>
    <h2>Source evidence appendix</h2><table><tr><th>ID</th><th>File row</th><th>Month</th><th>Department</th><th>Account</th><th>Scenario</th><th>Amount</th></tr>{appendix}</table></html>'''
