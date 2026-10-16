"""Strict input contract, integer-cent storage, immutable dataset versions."""
import csv
import hashlib
import io
import json
import os
import re
import sqlite3
import uuid
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from zipfile import ZipFile, BadZipFile

import pandas as pd

ROOT = Path(__file__).resolve().parent.parent
REQUIRED = ['period', 'department', 'account', 'category', 'scenario', 'amount']
ALIASES = {'month': 'period', 'date': 'period', 'dept': 'department', 'type': 'scenario', 'value': 'amount'}


def connection():
    path = Path(os.environ.get('FINANCEOS_DB', ROOT / 'data' / 'financeos.db'))
    path.parent.mkdir(parents=True, exist_ok=True)
    db = sqlite3.connect(path, timeout=20)
    db.row_factory = sqlite3.Row
    db.execute('PRAGMA foreign_keys=ON')
    return db


def init_db():
    with connection() as db:
        db.executescript('''
        CREATE TABLE IF NOT EXISTS datasets (
            id TEXT PRIMARY KEY, name TEXT NOT NULL, created TEXT NOT NULL,
            sha256 TEXT NOT NULL, synthetic INTEGER NOT NULL, validation TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS records (
            id INTEGER PRIMARY KEY, dataset_id TEXT NOT NULL REFERENCES datasets(id),
            source_row INTEGER NOT NULL, period TEXT NOT NULL, department TEXT NOT NULL,
            account TEXT NOT NULL, category TEXT NOT NULL, scenario TEXT NOT NULL,
            cents INTEGER NOT NULL, currency TEXT NOT NULL,
            UNIQUE(dataset_id,period,department,account,scenario));
        CREATE INDEX IF NOT EXISTS records_scope ON records(dataset_id,period,department);
        CREATE TABLE IF NOT EXISTS investigations (
            id TEXT PRIMARY KEY, dataset_id TEXT NOT NULL REFERENCES datasets(id), created TEXT NOT NULL, result TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS scenarios (
            id TEXT PRIMARY KEY, dataset_id TEXT NOT NULL REFERENCES datasets(id), created TEXT NOT NULL, result TEXT NOT NULL);
        ''')


def now():
    return datetime.now(timezone.utc).isoformat()


def parse_file(raw, filename):
    if len(raw) > 10 * 1024 * 1024:
        raise ValueError('File exceeds the 10 MB limit.')
    ext = Path(filename).suffix.lower()
    if ext == '.csv':
        header = next(csv.reader(io.StringIO(raw.decode('utf-8-sig'))), [])
        normalized = [ALIASES.get(c.strip().lower(), c.strip().lower()) for c in header]
        if len(set(normalized)) != len(normalized):
            raise ValueError('Duplicate column names after normalization.')
        df = pd.read_csv(io.BytesIO(raw), dtype=str, keep_default_na=False, encoding='utf-8-sig')
    elif ext == '.xlsx':
        try:
            with ZipFile(io.BytesIO(raw)) as z:
                if sum(i.file_size for i in z.infolist()) > 50 * 1024 * 1024:
                    raise ValueError('Expanded workbook exceeds 50 MB.')
        except BadZipFile as e:
            raise ValueError('This is not a valid XLSX workbook.') from e
        # data_only=False prevents cached formula values from silently entering the ledger.
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
        ws = wb.worksheets[0]
        values = list(ws.iter_rows(values_only=True))
        if not values:
            raise ValueError('The first worksheet is empty.')
        df = pd.DataFrame(values[1:], columns=values[0]).fillna('')
        wb.close()
    else:
        raise ValueError('Choose a CSV or XLSX file.')
    return validate_frame(df)


def validate_frame(df):
    if len(df) == 0 or len(df) > 50000:
        raise ValueError('Provide between 1 and 50,000 rows.')
    df = df.copy()
    df.columns = [ALIASES.get(str(c).strip().lower(), str(c).strip().lower()) for c in df.columns]
    if len(set(df.columns)) != len(df.columns):
        raise ValueError('Duplicate column names after normalization.')
    missing = set(REQUIRED) - set(df.columns)
    if missing:
        raise ValueError('Missing columns: ' + ', '.join(sorted(missing)))
    rows, errors, warnings, seen, mappings = [], [], [], {}, {}
    cleaned = 0
    for index, original in enumerate(df.to_dict('records'), 2):
        if all(str(v).strip() == '' for v in original.values()):
            warnings.append(f'Row {index}: removed empty row.')
            continue
        r = {k: str(original[k]).strip() for k in REQUIRED}
        if any(str(original[k]) != r[k] for k in REQUIRED):
            cleaned += 1
        try:
            period = original['period']
            if isinstance(period, (datetime, pd.Timestamp)):
                r['period'] = period.strftime('%Y-%m')
            if not re.fullmatch(r'20\d{2}-(0[1-9]|1[0-2])', r['period']):
                raise ValueError('period must be YYYY-MM (2000–2099)')
            for col in ('department', 'account'):
                if not r[col] or len(r[col]) > 100 or r[col][0] in '=+@':
                    raise ValueError(f'{col} must be a plain label, 1–100 characters')
            r['category'] = r['category'].lower()
            r['scenario'] = r['scenario'].lower()
            if r['category'] not in ('revenue', 'cogs', 'opex'):
                raise ValueError('category must be revenue, cogs or opex')
            if r['scenario'] not in ('actual', 'budget'):
                raise ValueError('scenario must be actual or budget')
            currency = str(original.get('currency', 'USD')).strip().upper()
            if currency != 'USD':
                raise ValueError('only USD is supported; convert currencies before upload')
            val = r['amount'].replace(',', '').replace('$', '')
            if val.startswith('(') and val.endswith(')'):
                val = '-' + val[1:-1]
            amount = Decimal(val)
            if not amount.is_finite() or abs(amount) > Decimal('100000000000'):
                raise ValueError('amount must be finite and within ±100 billion USD')
            if amount != amount.quantize(Decimal('.01')):
                raise ValueError('amount must have at most two decimal places')
            r['cents'] = int(amount * 100)
            # Signed corrections are supported; expenses normally positive.
            if amount < 0:
                warnings.append(f'Row {index}: negative amount retained as a signed correction.')
            key = tuple(r[k] for k in ('period', 'department', 'account', 'scenario'))
            mapping_key = r['account'].casefold()
            prior_category = mappings.setdefault(mapping_key, r['category'])
            if prior_category != r['category']:
                raise ValueError('account maps to inconsistent categories')
            if key in seen:
                if seen[key] == (r['cents'], r['category']):
                    warnings.append(f'Row {index}: exact duplicate removed.')
                    continue
                raise ValueError('conflicting duplicate key; aggregate transactions before upload')
            seen[key] = (r['cents'], r['category'])
            r.update(source_row=index, currency='USD')
            r.pop('amount')
            rows.append(r)
        except (ValueError, InvalidOperation) as e:
            errors.append({'row': index, 'message': str(e) or 'invalid numeric amount'})
    pair_keys = {}
    for r in rows:
        pair_keys.setdefault(tuple(r[k] for k in ('period','department','account')), set()).add(r['scenario'])
    unpaired = sum(len(s) != 2 for s in pair_keys.values())
    if unpaired:
        warnings.append(f'{unpaired} account/month/department keys lack an actual or budget counterpart. Variances are unavailable for these keys.')
    if cleaned:
        warnings.append(f'Trimmed whitespace in {cleaned} rows; normalized scenario/category casing and amount formatting.')
    if not rows and not errors:
        errors.append({'row': 0, 'message': 'No financial rows remain after cleaning.'})
    return rows, {'input_rows': len(df), 'accepted_rows': len(rows), 'errors': errors,
                  'warnings': warnings, 'unpaired_keys': unpaired, 'total_cents': sum(r['cents'] for r in rows),
                  'totals_by_scenario_cents': {s: sum(r['cents'] for r in rows if r['scenario']==s) for s in ('actual','budget')}}


def save_dataset(raw, filename, synthetic=False):
    rows, validation = parse_file(raw, filename)
    if validation['errors']:
        return None, validation
    dataset_id = uuid.uuid4().hex[:12]
    with connection() as db:
        db.execute('INSERT INTO datasets VALUES (?,?,?,?,?,?)',
                   (dataset_id, Path(filename).name, now(), hashlib.sha256(raw).hexdigest(), int(synthetic), json.dumps(validation)))
        db.executemany('''INSERT INTO records(dataset_id,source_row,period,department,account,category,scenario,cents,currency)
                          VALUES (:dataset_id,:source_row,:period,:department,:account,:category,:scenario,:cents,:currency)''',
                       [dict(r, dataset_id=dataset_id) for r in rows])
    return dataset_id, validation


def datasets():
    with connection() as db:
        return [dict(r) for r in db.execute('SELECT id,name,created,synthetic,sha256 FROM datasets ORDER BY created DESC')]


def metadata(dataset_id):
    with connection() as db:
        r = db.execute('SELECT * FROM datasets WHERE id=?', (dataset_id,)).fetchone()
    if not r:
        raise ValueError('Dataset not found.')
    d = dict(r)
    d['validation'] = json.loads(d['validation'])
    return d


def records(dataset_id, start=None, end=None, department=None):
    metadata(dataset_id)
    sql = 'SELECT * FROM records WHERE dataset_id=?'
    args = [dataset_id]
    for clause, value in [('period>=?',start),('period<=?',end),('department=?',department)]:
        if value:
            sql += ' AND ' + clause
            args.append(value)
    with connection() as db:
        return [dict(r) for r in db.execute(sql + ' ORDER BY period,department,account,scenario', args)]


def store_result(table, dataset_id, result):
    if table not in ('investigations', 'scenarios'):
        raise ValueError('Invalid result type.')
    result['id'] = uuid.uuid4().hex[:12]
    result['created'] = now()
    with connection() as db:
        db.execute(f'INSERT INTO {table} VALUES (?,?,?,?)', (result['id'], dataset_id, result['created'], json.dumps(result)))
    return result


def saved_result(table, result_id):
    if table not in ('investigations', 'scenarios'):
        raise ValueError('Invalid result type.')
    with connection() as db:
        r = db.execute(f'SELECT result FROM {table} WHERE id=?', (result_id,)).fetchone()
    if not r:
        raise ValueError('Saved result not found.')
    return json.loads(r[0])
